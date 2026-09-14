"""Excel 报表输出 (openpyxl)"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from xuanpin.paths import REPORTS_DIR

HEADERS = ["排名", "商品标题", "商品链接", "批发价低(¥)", "批发价高(¥)", "起订量",
           "销量线索", "重量(g,估)", "落地成本(¥)", "供货价上限(¥)",
           "保守利润(¥)", "乐观利润(¥)", "保守利润率", "风险标记", "采集时间", "零售价来源"]

# 保守利润 = 按批发价高值算; 乐观利润 = 按批发价低值算
COL_WIDTHS = [5, 42, 12, 11, 11, 8, 13, 10, 11, 13, 11, 11, 11, 30, 17, 22]
MONEY_COLS = [4, 5, 9, 10]
PROFIT_COLS = [11, 12]
PCT_COLS = [13]


def build(keyword, retail_price, rows, cfg):
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = REPORTS_DIR / f"选品报告_{keyword}_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "候选清单"
    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill

    rows_sorted = sorted(
        rows,
        key=lambda r: r["profit_low"] if r.get("profit_low") is not None else -1e9,
        reverse=True,
    )
    red = Font(color="C00000")
    for i, r in enumerate(rows_sorted, 1):
        ws.append([i, r.get("title"), r.get("url"), r.get("price_min"), r.get("price_max"),
                   r.get("moq"), r.get("sales_text"), r.get("weight_g"), r.get("landed"),
                   r.get("supply_cap"), r.get("profit_low"), r.get("profit_high"),
                   r.get("margin_low"), "; ".join(r.get("flags", [])), r.get("captured_at"),
                   r.get("retail_source", "人工")])
        ri = ws.max_row
        link = ws.cell(row=ri, column=3)
        if r.get("url"):
            link.hyperlink = r["url"]
            link.font = Font(color="0563C1", underline="single")
        for col in MONEY_COLS:
            ws.cell(row=ri, column=col).number_format = "0.00"
        for col in PROFIT_COLS:
            c = ws.cell(row=ri, column=col)
            c.number_format = "0.00"
            if isinstance(c.value, (int, float)) and c.value < 0:
                c.font = red
        for col in PCT_COLS:
            c = ws.cell(row=ri, column=col)
            c.number_format = "0.0%"
            if isinstance(c.value, (int, float)) and c.value < 0:
                c.font = red
        if r.get("flags"):
            ws.cell(row=ri, column=14).font = red

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ps = wb.create_sheet("运行参数")
    ps.append(["参数", "值"])
    for k, v in [
        ("关键词", keyword),
        ("目标零售价(¥)", retail_price),
        ("平台分成比例", cfg["fees"]["platform_take"]),
        ("包装费(¥/件)", cfg["costing"]["packaging_fee"]),
        ("国内送仓(¥/件)", cfg["costing"]["inbound_ship"]),
        ("默认重量(g)", cfg["costing"]["default_weight_g"]),
        ("候选商品数", len(rows_sorted)),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("口径说明", "供货价上限=零售价×(1-分成); 落地成本=采购价+包装+送仓(+超重); "
                  "保守=批发价取高值, 乐观=取低值"),
    ]:
        ps.append([k, v])
    ps.column_dimensions["A"].width = 16
    ps.column_dimensions["B"].width = 100

    wb.save(path)
    return path


SUMMARY_HEADERS = ["关键词", "对标零售价(¥)", "排名", "商品标题", "批发价低(¥)", "批发价高(¥)",
                   "起订量", "销量线索", "重量(g,估)", "落地成本(¥)", "供货价上限(¥)",
                   "保守利润(¥)", "乐观利润(¥)", "保守利润率", "风险标记", "采集时间", "零售价来源"]
SUMMARY_WIDTHS = [12, 12, 5, 40, 10, 10, 7, 12, 9, 10, 12, 10, 10, 10, 26, 16, 22]


def build_summary(results, cfg):
    """批量工作流汇总表: results = [{keyword, retail, rows, per_report}, ...]"""
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = REPORTS_DIR / f"汇总选品报告_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(SUMMARY_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7B3F00")

    flat = []
    for r in results:
        for row in r["rows"]:
            flat.append((r["keyword"], r["retail"], row))
    flat.sort(key=lambda x: x[2]["profit_low"], reverse=True)

    red = Font(color="C00000")
    for i, (kw, retail, r) in enumerate(flat, 1):
        ws.append([kw, retail, i, r.get("title"), r.get("price_min"), r.get("price_max"),
                   r.get("moq"), r.get("sales_text"), r.get("weight_g"), r.get("landed"),
                   r.get("supply_cap"), r.get("profit_low"), r.get("profit_high"),
                   r.get("margin_low"), "; ".join(r.get("flags", [])), r.get("captured_at"),
                   r.get("retail_source", "人工")])
        ri = ws.max_row
        for col in (5, 6, 10, 11):
            ws.cell(row=ri, column=col).number_format = "0.00"
        for col in (12, 13):
            c = ws.cell(row=ri, column=col)
            c.number_format = "0.00"
            if isinstance(c.value, (int, float)) and c.value < 0:
                c.font = red
        for col in (14,):
            c = ws.cell(row=ri, column=col)
            c.number_format = "0.0%"
            if isinstance(c.value, (int, float)) and c.value < 0:
                c.font = red
        if r.get("flags"):
            ws.cell(row=ri, column=15).font = red

    for i, w in enumerate(SUMMARY_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ms = wb.create_sheet("本次清单")
    ms.append(["关键词", "对标零售价(¥)", "候选数", "最高保守利润(¥)", "单品报告路径"])
    for r in results:
        best = max((x["profit_low"] for x in r["rows"]), default=None)
        ms.append([r["keyword"], r["retail"], len(r["rows"]), best,
                   str(r["per_report"]) if r.get("per_report") else ("-" if not r["rows"] else "demo模式不生成")])
    for i, w in enumerate([12, 14, 8, 16, 70], 1):
        ms.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    return path


SUPPLY_HEADERS = ["状态", "建议", "原因", "商品标题", "录入拿货价(¥)", "采购现价(¥)", "价格变化",
                  "供货价(¥)", "录入时利润(¥)", "当前利润(¥)", "当前利润率", "重量(g)",
                  "最近采集", "商品链接"]
SUPPLY_WIDTHS = [9, 15, 34, 36, 12, 11, 9, 10, 12, 11, 10, 8, 11, 14]
SUGGESTION_COLOR = {"建议下架(停供)": "C00000", "建议换源/停供": "C00000",
                    "谨慎持有": "B26A00", "信息过期": "7030A0", "建议加量/续报": "1A7F45"}


def build_supply_report(rows, cfg):
    """供货监控报告: 每个在供商品的最新采购价、利润和上下架建议"""
    from xuanpin.supply import _thresholds
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = REPORTS_DIR / f"供货监控报告_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "供货监控"
    ws.append(SUPPLY_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5E4E")

    red = Font(color="C00000")
    for r in sorted(rows, key=lambda x: x["suggestion"]):
        chg = r.get("price_change")
        ws.append(["在供" if r.get("active") else "已停供", r["suggestion"],
                   r.get("reasons") or "", r.get("title"), r.get("purchase_price"),
                   r.get("purchase_now"), chg, r.get("supply_price"),
                   r.get("profit_entry"), r.get("profit_now"), r.get("ratio_now"),
                   r.get("weight_g"), (r.get("captured_at") or "")[:10], r.get("url")])
        ri = ws.max_row
        sug = ws.cell(row=ri, column=2)
        color = SUGGESTION_COLOR.get(r["suggestion"])
        if color:
            sug.font = Font(color=color, bold=True)
        for col in (5, 6, 8, 9, 10):
            ws.cell(row=ri, column=col).number_format = "0.00"
        c7 = ws.cell(row=ri, column=7)
        c7.number_format = "+0.0%;-0.0%"
        if isinstance(chg, (int, float)) and chg > 0:
            c7.font = red
        c11 = ws.cell(row=ri, column=11)
        c11.number_format = "0.0%"
        if isinstance(c11.value, (int, float)) and c11.value < 0.10:
            c11.font = red
        link = ws.cell(row=ri, column=14)
        if r.get("url"):
            link.hyperlink = r["url"]
            link.font = Font(color="0563C1", underline="single")

    for i, w in enumerate(SUPPLY_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ts2 = wb.create_sheet("口径说明")
    t = _thresholds(cfg)
    for k, v in [("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                 ("下架线", f"利润率≤{t['min_ratio_stop']*100:.0f}% 或亏本 → 建议下架(停供)"),
                 ("谨慎线", f"利润率<{t['min_ratio_warn']*100:.0f}% → 核价空间不足"),
                 ("加量线", f"利润率≥{t['boost_ratio']*100:.0f}%且成本降{t['boost_drop']*100:.0f}% → 建议加量/续报"),
                 ("断供判定", f"超过{t['stale_days']}天没采集到货源价格 → 提示换源"),
                 ("当前利润", "供货价 - (最新采购价+包装+送仓+超重)"),
                 ("提示", "Temu全托管由平台控制上架, 建议需人工到商家后台执行: 停止供货/调整报价")]:
        ts2.append([k, v])
    ts2.column_dimensions["A"].width = 14
    ts2.column_dimensions["B"].width = 90

    wb.save(path)
    return path
